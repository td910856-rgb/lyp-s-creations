"""把上传的简历文件转成纯文本。

支持：PDF（文字版）、DOCX、TXT / MD。
不支持：扫描件 PDF（纯图片，需要 OCR）、旧版 .doc。
所有失败都抛 ParseError，message 是能直接给用户看的中文提示。
"""

from __future__ import annotations

import re
from pathlib import Path

# 允许上传的扩展名
SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".txt", ".md"}

# 提取出的文本少于这个长度，基本可以判断是没解析成功
MIN_TEXT_CHARS = 50


class ParseError(Exception):
    """文件解析失败。message 会原样返回给前端展示。"""


def extract_text(path: Path) -> str:
    """读取文件并返回清洗后的纯文本。"""
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        text = _from_pdf(path)
    elif suffix == ".docx":
        text = _from_docx(path)
    elif suffix in {".txt", ".md"}:
        text = _from_plain_text(path)
    elif suffix == ".doc":
        raise ParseError("旧版 .doc 格式读不了，请在 Word 里「另存为」成 .docx 再上传。")
    else:
        raise ParseError(f"暂不支持 {suffix or '这种'} 格式，请上传 PDF 或 DOCX 文件。")

    text = _clean(text)

    if len(text) < MIN_TEXT_CHARS:
        raise ParseError(
            "这份文件里几乎读不出文字。如果它是扫描件或图片版 PDF，"
            "请换一份文字版 PDF，或直接把简历内容复制成 .txt 上传。"
        )
    return text


def _from_pdf(path: Path) -> str:
    try:
        import pdfplumber
    except ImportError as exc:  # pragma: no cover
        raise ParseError("缺少 pdfplumber 依赖，请先执行 pip install -r requirements.txt") from exc

    pages: list[str] = []
    try:
        with pdfplumber.open(path) as pdf:
            if not pdf.pages:
                raise ParseError("这个 PDF 是空的，没有任何页面。")
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
    except ParseError:
        raise
    except Exception as exc:
        raise ParseError(f"PDF 读取失败：{exc}。可以试试把文件另存为后重新上传。") from exc

    return "\n".join(pages)


def _from_docx(path: Path) -> str:
    try:
        from docx import Document
    except ImportError as exc:  # pragma: no cover
        raise ParseError("缺少 python-docx 依赖，请先执行 pip install -r requirements.txt") from exc

    try:
        document = Document(str(path))
    except Exception as exc:
        raise ParseError(f"Word 文件读取失败：{exc}。请确认文件没有损坏。") from exc

    blocks: list[str] = [p.text for p in document.paragraphs]

    # 简历经常用表格排版，表格里的文字也要读出来
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            line = " | ".join(cell for cell in cells if cell)
            if line:
                blocks.append(line)

    return "\n".join(blocks)


def _from_plain_text(path: Path) -> str:
    raw = path.read_bytes()
    for encoding in ("utf-8-sig", "utf-8", "gb18030"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError("文本编码无法识别，请把文件另存为 UTF-8 编码后再上传。")


def _clean(text: str) -> str:
    """清洗文本：统一换行、去掉多余空行、压缩重复空格。"""
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    lines = [re.sub(r"[ \t\u3000]+", " ", line).strip() for line in text.split("\n")]
    lines = [line for line in lines if line]
    return "\n".join(lines).strip()

